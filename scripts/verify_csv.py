"""Verifica en profundidad que los CSV en data/csv/ contienen toda la
información de los .xlsx en data/raw/Hackaton/.

Compara celda a celda lo que openpyxl lee del Excel (valores cacheados de
fórmulas, fechas, números) contra lo que hay en el CSV. Detecta:
- Hojas faltantes o sobrantes (incluidas hojas ocultas).
- Diferencias de dimensiones.
- Diferencias de contenido por celda.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = REPO_ROOT / "data" / "raw" / "Hackaton"
OUTPUT_DIR = REPO_ROOT / "data" / "csv"


def slugify(name: str) -> str:
    name = name.strip()
    name = re.sub(r"\s+", "_", name)
    name = re.sub(r"[^A-Za-z0-9._-]", "", name)
    return name or "sheet"


def cell_to_str(cell: Cell) -> str:
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v:
            return ""
        if v.is_integer() and abs(v) < 1e16:
            return str(int(v))
        return repr(v)
    if isinstance(v, dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0 and v.microsecond == 0:
            return v.date().isoformat()
        return v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    return str(v)


def excel_sheet_to_matrix(xlsx_path: Path, sheet_name: str) -> list[list[str]]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        max_col = ws.max_column or 0
        matrix: list[list[str]] = []
        for row in ws.iter_rows(values_only=False):
            values = [cell_to_str(c) for c in row]
            if len(values) < max_col:
                values.extend([""] * (max_col - len(values)))
            matrix.append(values)
        # Normaliza: si la última fila/col son todas vacías, las quitamos
        while matrix and all(c == "" for c in matrix[-1]):
            matrix.pop()
        if matrix:
            ncols = max(len(r) for r in matrix)
            matrix = [r + [""] * (ncols - len(r)) for r in matrix]
            while ncols > 0 and all(r[ncols - 1] == "" for r in matrix):
                for r in matrix:
                    r.pop()
                ncols -= 1
        return matrix
    finally:
        wb.close()


def read_csv_matrix(csv_path: Path) -> list[list[str]]:
    with csv_path.open("r", encoding="utf-8", newline="") as fh:
        rows = list(csv.reader(fh))
    while rows and all(c == "" for c in rows[-1]):
        rows.pop()
    if rows:
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        while ncols > 0 and all(r[ncols - 1] == "" for r in rows):
            for r in rows:
                r.pop()
            ncols -= 1
    return rows


def compare_matrices(
    label: str, excel: list[list[str]], csvm: list[list[str]]
) -> list[str]:
    errors: list[str] = []
    if len(excel) != len(csvm):
        errors.append(
            f"{label}: filas Excel={len(excel)} vs CSV={len(csvm)}"
        )
    excel_cols = max((len(r) for r in excel), default=0)
    csv_cols = max((len(r) for r in csvm), default=0)
    if excel_cols != csv_cols:
        errors.append(f"{label}: columnas Excel={excel_cols} vs CSV={csv_cols}")
    rows = min(len(excel), len(csvm))
    cols = min(excel_cols, csv_cols)
    diffs = 0
    examples: list[str] = []
    for i in range(rows):
        for j in range(cols):
            a = excel[i][j] if j < len(excel[i]) else ""
            b = csvm[i][j] if j < len(csvm[i]) else ""
            if a != b:
                diffs += 1
                if len(examples) < 5:
                    examples.append(
                        f"fila={i} col={j}: Excel={a!r} vs CSV={b!r}"
                    )
    if diffs:
        errors.append(f"{label}: {diffs} celdas distintas. Ejemplos: {examples}")
    return errors


def main() -> int:
    if not SOURCE_DIR.exists():
        print(f"No existe {SOURCE_DIR}")
        return 1

    xlsx_files = sorted(p for p in SOURCE_DIR.rglob("*") if p.suffix.lower() == ".xlsx")
    if not xlsx_files:
        print("No hay archivos xlsx que verificar.")
        return 0

    total_errors: list[str] = []
    print(f"Verificando {len(xlsx_files)} archivos Excel...\n")

    for xlsx_path in xlsx_files:
        rel = xlsx_path.relative_to(SOURCE_DIR)
        print(f"=== {rel} ===")
        workbook_dir = OUTPUT_DIR / slugify(xlsx_path.stem)

        if not workbook_dir.exists():
            msg = f"FALTA carpeta de salida {workbook_dir}"
            print(f"  ✗ {msg}")
            total_errors.append(f"{rel}: {msg}")
            continue

        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        sheet_names = list(wb.sheetnames)
        hidden = [s for s in sheet_names if wb[s].sheet_state != "visible"]
        wb.close()
        print(f"  Hojas en Excel ({len(sheet_names)}): {sheet_names}")
        if hidden:
            print(f"  ⚠ Ocultas: {hidden}")

        existing_csvs = sorted(p.name for p in workbook_dir.glob("*.csv"))
        expected_csvs = sorted(f"{slugify(s)}.csv" for s in sheet_names)
        unexpected = set(existing_csvs) - set(expected_csvs)
        missing = set(expected_csvs) - set(existing_csvs)
        if missing:
            msg = f"FALTAN CSVs: {sorted(missing)}"
            print(f"  ✗ {msg}")
            total_errors.append(f"{rel}: {msg}")
        if unexpected:
            msg = f"CSVs inesperados: {sorted(unexpected)}"
            print(f"  ✗ {msg}")
            total_errors.append(f"{rel}: {msg}")

        for sheet_name in sheet_names:
            csv_path = workbook_dir / f"{slugify(sheet_name)}.csv"
            if not csv_path.exists():
                continue
            try:
                excel_m = excel_sheet_to_matrix(xlsx_path, sheet_name)
                csv_m = read_csv_matrix(csv_path)
            except Exception as exc:  # noqa: BLE001
                msg = f"Error leyendo '{sheet_name}': {exc}"
                print(f"  ✗ {msg}")
                total_errors.append(f"{rel} :: {msg}")
                continue
            label = f"'{sheet_name}'"
            errors = compare_matrices(label, excel_m, csv_m)
            if errors:
                for e in errors:
                    print(f"  ✗ {e}")
                    total_errors.append(f"{rel} :: {e}")
            else:
                ex_r = len(excel_m)
                ex_c = max((len(r) for r in excel_m), default=0)
                cs_r = len(csv_m)
                cs_c = max((len(r) for r in csv_m), default=0)
                print(
                    f"  ✓ {label} OK "
                    f"(Excel: {ex_r}f x {ex_c}c, CSV: {cs_r}f x {cs_c}c)"
                )
        print()

    print("=" * 60)
    if total_errors:
        print(f"Se encontraron {len(total_errors)} problemas:")
        for e in total_errors:
            print(f" - {e}")
        return 2

    print("Todo OK: cada hoja de cada Excel está representada y los datos coinciden celda a celda.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
