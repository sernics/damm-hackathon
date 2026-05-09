"""Convierte todos los .xlsx/.XLSX dentro de data/raw/Hackaton a CSV.

Lee con openpyxl celda a celda para preservar los valores exactos
(enteros sin .0, fechas en formato ISO, fórmulas evaluadas a su valor).

Cada archivo Excel puede tener varias hojas (incluidas ocultas). Cada hoja
se exporta a un CSV independiente dentro de una subcarpeta nombrada como
el archivo original.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell

REPO_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = REPO_ROOT / "data" / "raw" / "Hackaton"
OUTPUT_DIR = REPO_ROOT / "data" / "csv"


def slugify(name: str) -> str:
    """Normaliza nombres para usarlos como nombres de archivo/carpeta."""
    name = name.strip()
    name = re.sub(r"\s+", "_", name)
    name = re.sub(r"[^A-Za-z0-9._-]", "", name)
    return name or "sheet"


def cell_to_str(cell: Cell) -> str:
    """Convierte una celda openpyxl a su representación textual exacta."""
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v:  # NaN
            return ""
        if v.is_integer() and abs(v) < 1e16:
            return str(int(v))
        return repr(v)
    if isinstance(v, dt.datetime):
        if v.hour == 0 and v.minute == 0 and v.second == 0 and v.microsecond == 0:
            return v.date().isoformat()
        return v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    return str(v)


def convert_workbook(xlsx_path: Path, out_root: Path) -> list[Path]:
    """Convierte un workbook a uno o varios CSV (uno por hoja)."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    workbook_dir = out_root / slugify(xlsx_path.stem)
    workbook_dir.mkdir(parents=True, exist_ok=True)

    written: list[Path] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        csv_path = workbook_dir / f"{slugify(sheet_name)}.csv"

        with csv_path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh, quoting=csv.QUOTE_MINIMAL, lineterminator="\n")

            max_col = ws.max_column or 0
            for row in ws.iter_rows(values_only=False):
                values = [cell_to_str(c) for c in row]
                if len(values) < max_col:
                    values.extend([""] * (max_col - len(values)))
                writer.writerow(values)

        written.append(csv_path)
    wb.close()
    return written


def main() -> None:
    if not SOURCE_DIR.exists():
        raise SystemExit(f"No existe el directorio de origen: {SOURCE_DIR}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    xlsx_files = sorted(
        p for p in SOURCE_DIR.rglob("*") if p.suffix.lower() == ".xlsx"
    )
    if not xlsx_files:
        print(f"No se encontraron archivos .xlsx en {SOURCE_DIR}")
        return

    print(f"Encontrados {len(xlsx_files)} archivos Excel.")
    for xlsx_path in xlsx_files:
        rel = xlsx_path.relative_to(SOURCE_DIR)
        print(f"\n→ Procesando: {rel}")
        try:
            written = convert_workbook(xlsx_path, OUTPUT_DIR)
        except Exception as exc:  # noqa: BLE001
            print(f"  ✗ Error procesando {rel}: {exc}")
            continue
        for csv_path in written:
            print(f"  ✓ {csv_path.relative_to(REPO_ROOT)}")


if __name__ == "__main__":
    main()
